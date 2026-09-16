import argparse
import re
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, select

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

from app.database import Base, SessionLocal, engine  # noqa: E402
from app.models import Food, FoodNutrient, ProductStandardMapping, StandardFood  # noqa: E402
from app.services.mfds import MfdsClient  # noqa: E402

NUTRIENT_COLUMNS = {
    "kcal": (6, "kcal"), "protein": (8, "g"), "calcium": (22, "mg"),
    "iron": (23, "mg"), "vitaminA": (34, "µgRAE"), "vitaminC": (51, "mg"),
    "sodium": (27, "mg"),
}

DEMO_PRODUCTS = [
    {"cu_id": "5228", "serving_amount": 75, "serving_unit": "g", "standard_code": None,
     "mfds_query": "바나나킥", "mfds_maker": "농심"},
    {"cu_id": "6979", "serving_amount": 240, "serving_unit": "g", "standard_code": "1461"},
    {"cu_id": "8712", "serving_amount": 100, "serving_unit": "g", "standard_code": "1650"},
    {"cu_id": "6216", "serving_amount": 190, "serving_unit": "ml", "standard_code": "571"},
    {"cu_id": "17620", "serving_amount": 180, "serving_unit": "g", "standard_code": None},
]


def existing_file(value: str) -> Path:
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise argparse.ArgumentTypeError(f"파일을 찾을 수 없습니다: {path}")
    return path


def number(value: object) -> float | None:
    if value is None or str(value).strip() in {"", "-", "N/A", "Tr"}:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def extract_barcode(url: str | None) -> str | None:
    match = re.search(r"(\d{13})(?:_\d+)?\.(?:jpg|png)$", url or "", re.IGNORECASE)
    return match.group(1) if match else None


def load_cu_rows(path: Path) -> dict[str, dict]:
    ws = load_workbook(path, read_only=True, data_only=True)["CU 식품·음료"]
    wanted = {item["cu_id"] for item in DEMO_PRODUCTS}
    result = {}
    for row in ws.iter_rows(min_row=8, values_only=True):
        cu_id = str(row[5] or "")
        if cu_id in wanted:
            result[cu_id] = {
                "category": row[0], "name": row[1],
                "price": int(row[2]) if row[2] is not None else None,
                "image_url": row[7], "barcode": extract_barcode(row[7]),
            }
    return result


def load_standard_rows(path: Path) -> dict[str, dict]:
    ws = load_workbook(path, read_only=True, data_only=True)["국가표준식품성분 Database 10.4"]
    wanted = {item["standard_code"] for item in DEMO_PRODUCTS if item.get("standard_code")}
    result = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        code = str(row[0] or "")
        if code in wanted:
            result[code] = {
                "name": str(row[3]), "food_group": str(row[2]),
                "nutrients": {key: number(row[column - 1]) for key, (column, _) in NUTRIENT_COLUMNS.items()},
            }
    return result


def mfds_match(config: dict):
    if not config.get("mfds_query"):
        return None
    try:
        matches = MfdsClient().search(config["mfds_query"], config.get("mfds_maker"))
    except Exception as error:
        print(f"식약처 조회를 건너뜁니다: {type(error).__name__}")
        return None
    return next((item for item in matches if item.item_report_no == "199104611019"), None) or (matches[0] if matches else None)


def seed(cu_path: Path, standard_path: Path) -> None:
    cu_rows = load_cu_rows(cu_path)
    standard_rows = load_standard_rows(standard_path)
    if len(cu_rows) != len(DEMO_PRODUCTS):
        raise RuntimeError("선택한 CU 테스트 상품을 모두 찾지 못했습니다.")

    Base.metadata.create_all(bind=engine)
    with SessionLocal() as db:
        demo_ids = [f"cu:{item['cu_id']}" for item in DEMO_PRODUCTS]
        db.execute(delete(ProductStandardMapping).where(ProductStandardMapping.food_id.in_(demo_ids)))
        db.execute(delete(FoodNutrient).where(FoodNutrient.food_id.in_(demo_ids)))
        db.execute(delete(Food).where(Food.id.in_(demo_ids)))

        for config in DEMO_PRODUCTS:
            source = cu_rows[config["cu_id"]]
            mfds = mfds_match(config)
            standard_code = config.get("standard_code")
            standard = standard_rows.get(standard_code)
            food = Food(
                id=f"cu:{config['cu_id']}", source_type="CU_PRODUCT", name=source["name"],
                brand=source["name"].split(")", 1)[0] if ")" in source["name"] else "CU",
                category=source["category"], price=source["price"], barcode=source["barcode"],
                image_url=source["image_url"],
                item_report_no=mfds.item_report_no if mfds else f"TEST-CU-{config['cu_id']}",
                serving_amount=Decimal(str(config["serving_amount"])),
                serving_unit=config["serving_unit"], count_unit="개",
                serving_label=f"{config['serving_amount']}{config['serving_unit']}",
            )
            db.add(food)

            if standard:
                if db.get(StandardFood, standard_code) is None:
                    db.add(StandardFood(food_code=standard_code, name=standard["name"],
                                        food_group=standard["food_group"], nutrients_per_100g=standard["nutrients"]))
                db.add(ProductStandardMapping(food_id=food.id, standard_food_code=standard_code,
                                               match_method="DEMO_MANUAL", confidence=Decimal("0.8500"),
                                               review_status="TEST"))

            for nutrient_key, (_, unit) in NUTRIENT_COLUMNS.items():
                value, quality, nutrient_source, source_ref = None, "MISSING", "NONE", None
                if mfds and mfds.nutrients_per_basis.get(nutrient_key) is not None:
                    value = mfds.nutrients_per_basis[nutrient_key] * config["serving_amount"] / mfds.basis_amount
                    quality, nutrient_source, source_ref = "CONFIRMED", "MFDS", mfds.item_report_no
                elif standard and standard["nutrients"].get(nutrient_key) is not None:
                    value = standard["nutrients"][nutrient_key] * config["serving_amount"] / 100
                    quality, nutrient_source, source_ref = "ESTIMATED", "NATIONAL_STANDARD", standard_code
                db.add(FoodNutrient(food_id=food.id, nutrient_key=nutrient_key,
                                    value=Decimal(str(round(value, 4))) if value is not None else None,
                                    unit=unit, quality=quality, source=nutrient_source, source_ref=source_ref))

        db.commit()
        print(f"CU 테스트 상품 {len(DEMO_PRODUCTS)}개를 적재했습니다.")
        for food in db.scalars(select(Food).where(Food.id.in_(demo_ids)).order_by(Food.id)):
            print(f"- {food.name}: {food.item_report_no}")


def main() -> None:
    parser = argparse.ArgumentParser(description="CU 테스트 상품과 영양 데이터를 적재합니다.")
    parser.add_argument("--cu", required=True, type=existing_file)
    parser.add_argument("--standard", required=True, type=existing_file)
    args = parser.parse_args()
    seed(args.cu, args.standard)


if __name__ == "__main__":
    main()
